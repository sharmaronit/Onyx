from pptx import Presentation
import glob
p=glob.glob(r'C:\Users\Ronit Sharma\Downloads\*.pptx')[0]
prs=Presentation(p)
print('file',p,'slides',len(prs.slides),'size',prs.slide_width,prs.slide_height)
for i,s in enumerate(prs.slides,1):
 print('\nSLIDE',i)
 for sh in s.shapes:
  if hasattr(sh,'text') and sh.text.strip(): print(repr(sh.text[:500]))
